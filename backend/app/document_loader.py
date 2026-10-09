"""Extract text from supported document types, keeping page labels when possible."""
from io import BytesIO
from pathlib import Path
import pymupdf as fitz
from docx import Document
from PIL import Image
import pytesseract

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".webp"}


def extract_document(filename: str, data: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or 'no extension'}")

    pages: list[dict] = []
    if ext == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as pdf:
            for index, page in enumerate(pdf):
                text = page.get_text("text").strip()
                # If a page has little embedded text, treat it as a scanned page and OCR it.
                if len(text) < 30:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = Image.open(BytesIO(pix.tobytes("png")))
                    ocr_text = pytesseract.image_to_string(image).strip()
                    if len(ocr_text) > len(text):
                        text = ocr_text
                if text:
                    pages.append({"page": index + 1, "text": text, "kind": "page"})
    elif ext == ".docx":
        doc = Document(BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for i, paragraph in enumerate(paragraphs):
            pages.append({"page": None, "text": paragraph, "kind": "paragraph", "index": i + 1})
        # Also include tables as readable row text. Table coordinates/pages are not available here.
        for table_index, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
            if rows:
                pages.append({"page": None, "text": f"Table {table_index}:\n" + "\n".join(rows), "kind": "table"})
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        image = Image.open(BytesIO(data)).convert("RGB")
        text = pytesseract.image_to_string(image).strip()
        pages.append({"page": 1, "text": text, "kind": "image_ocr"})
    elif ext == ".txt":
        pages.append({"page": None, "text": data.decode("utf-8", errors="replace"), "kind": "text"})

    if not any(p["text"].strip() for p in pages):
        raise ValueError("No text could be extracted. The file may be empty or require a different OCR tool.")
    return pages
